"""
Import de fiches d'inscription (photos / PDF) -> liste Excel.

Principe : on donne à Gemini la photo (ou le PDF) d'une fiche + la liste des colonnes
voulues (Nom, Prénom, Sexe, Date de naissance...). Il lit l'écriture manuscrite et renvoie
une ligne par fiche. Le code vérifie ensuite les dates, les téléphones, le sexe, retire les
doublons et écrit le tout dans un fichier .xlsx (nouveau, ou à la suite d'une feuille existante).

Aucune dépendance à Jarvis : main2.py appelle configurer(...) puis extraire_fiches / ecrire_xlsx.
"""
import io
import json
import re
import unicodedata
from copy import copy
from datetime import datetime

COLONNES_DEFAUT = [
    "Nom", "Prénom ( s )", "Sexe", "Date de naissance", "Lieu de naissance",
    "Téléphone de parents", "Nationalité", "Npi",
]

TYPES_ACCEPTES = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".webp": "image/webp", ".heic": "image/heic", ".heif": "image/heif",
    ".pdf": "application/pdf",
}

_MOIS = {
    "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6, "juillet": 7,
    "aout": 8, "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12,
}

_CTX = {"client": None, "types": None, "modeles": ["gemini-2.5-flash"]}


def configurer(client, types, modeles=None):
    """Branche le client Gemini de Jarvis (appelé une fois par main2.py)."""
    _CTX["client"] = client
    _CTX["types"] = types
    if modeles:
        _CTX["modeles"] = list(modeles)


# ──────────────────────────── utilitaires texte ────────────────────────────

def _norm(s):
    """Minuscules sans accents ni ponctuation : 'Prénom ( s )' -> 'prenom s'."""
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _vide(v):
    return v is None or str(v).strip() == "" or str(v).strip().lower() in ("null", "none", "n/a", "-")


def normaliser_date(valeur):
    """'30 janvier 2023', '7/2/23', '2023-01-30' -> ('30/01/2023', None) ou (texte, 'raison')."""
    if _vide(valeur):
        return "", None
    brut = str(valeur).strip()
    t = _norm(brut)
    j = m = a = None
    r = re.match(r"^(\d{1,2}) (\d{1,2}) (\d{2,4})$", t)
    if r:
        j, m, a = int(r[1]), int(r[2]), int(r[3])
    else:
        r = re.match(r"^(\d{4}) (\d{1,2}) (\d{1,2})$", t)
        if r:
            a, m, j = int(r[1]), int(r[2]), int(r[3])
        else:
            r = re.match(r"^(\d{1,2}) (?:er )?([a-z]+) (\d{2,4})$", t)
            if r and r[2] in _MOIS:
                j, m, a = int(r[1]), _MOIS[r[2]], int(r[3])
    if j is None:
        return brut, "date illisible ou format inconnu"
    if a < 100:
        a += 2000 if a < 50 else 1900
    try:
        datetime(a, m, j)
    except ValueError:
        return f"{j:02d}/{m:02d}/{a}", "date impossible (ce jour n'existe pas dans ce mois), à vérifier"
    return f"{j:02d}/{m:02d}/{a}", None


def normaliser_telephones(valeur):
    """Garde les numéros en chiffres, séparés par '/'. 8 chiffres -> préfixe 01 (Bénin)."""
    if _vide(valeur):
        return "", None
    morceaux = re.split(r"[/;,|]|\bet\b", str(valeur))
    nums = []
    for mo in morceaux:
        d = re.sub(r"\D", "", mo)
        if not d:
            continue
        if len(d) == 8:
            d = "01" + d
        if d not in nums:
            nums.append(d)
    if not nums:
        return str(valeur).strip(), "numéro illisible"
    doute = None
    if any(len(n) != 10 for n in nums):
        doute = "numéro de longueur inhabituelle, à vérifier"
    return "/".join(nums), doute


def normaliser_sexe(valeur):
    if _vide(valeur):
        return "", None
    t = _norm(valeur)
    if t in ("m", "masculin", "garcon", "g", "h", "homme"):
        return "M", None
    if t in ("f", "feminin", "fille"):
        return "F", None
    return str(valeur).strip(), "sexe non reconnu"


def _role(colonne):
    """Rôle d'une colonne d'après son nom (pour appliquer les bonnes vérifications)."""
    n = _norm(colonne)
    if n.startswith("sexe") or n == "genre":
        return "sexe"
    if "date" in n and "naiss" in n:
        return "date"
    if "tel" in n or "phone" in n or "contact" in n:
        return "tel"
    return None


# ───────────────────────────── lecture par Gemini ─────────────────────────────

def _prompt(colonnes, fixes, consignes):
    cols = "\n".join(f'  - "{c}"' for c in colonnes)
    fixes_txt = ""
    if fixes:
        fixes_txt = "\nCes colonnes sont remplies automatiquement, tu les laisses vides : " + \
                    ", ".join(f'"{k}"' for k in fixes) + "."
    extra = f"\nConsignes de l'utilisateur : {consignes.strip()}" if consignes and consignes.strip() else ""
    return f"""Tu lis des fiches de renseignement / d'inscription d'élèves (photo ou PDF, écriture manuscrite).
Chaque fiche concerne UN enfant. Une même image ou un même PDF peut contenir plusieurs fiches : renvoie une ligne par enfant.

Colonnes à remplir (utilise EXACTEMENT ces noms comme clés JSON) :
{cols}{fixes_txt}

Règles :
- "Nom" = nom de famille de l'enfant, EN MAJUSCULES. Prénom(s) = tous les prénoms, avec majuscule initiale. Si le nom et les prénoms sont écrits ensemble, sépare-les (le nom de famille est en général celui du père).
- Sexe : lis la case cochée (Garçon -> "M", Fille -> "F"). Si les deux cases ou aucune sont cochées, laisse vide et signale-le dans "_doutes".
- Date de naissance au format JJ/MM/AAAA. Lieu de naissance : seulement ce qui est écrit sur la fiche (souvent après la date). Laisse vide s'il n'est pas écrit.
- Téléphone des parents : numéro du père puis numéro de la mère, en chiffres sans espace, séparés par "/". Un seul numéro s'il n'y en a qu'un. Si le numéro d'un parent est écrit à un mauvais endroit de la fiche, utilise-le quand même.
- N'invente JAMAIS une valeur : si un champ est absent ou illisible, mets "" (chaîne vide).
- Dans "_doutes", liste en français chaque valeur incertaine, floue ou bizarre, sous la forme "Colonne: raison" (ex. "Date de naissance: le jour est peu lisible").{extra}

Réponds UNIQUEMENT par un tableau JSON, sans texte autour. Chaque élément est un objet avec les colonnes ci-dessus et une clé "_doutes" (tableau de textes, vide s'il n'y a aucun doute)."""


def _parse_json(texte):
    t = (texte or "").strip()
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t, flags=re.S).strip()
    try:
        data = json.loads(t)
    except Exception:
        m = re.search(r"\[.*\]", t, flags=re.S)
        if not m:
            raise ValueError("réponse non JSON")
        data = json.loads(m.group(0))
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list):
                data = v
                break
        else:
            data = [data]
    return [d for d in data if isinstance(d, dict)]


def _lire_un_fichier(nom, contenu, mime, colonnes, fixes, consignes):
    client, types = _CTX["client"], _CTX["types"]
    if client is None or types is None:
        raise RuntimeError("Gemini n'est pas configuré (clé GEMINI_API_KEY manquante).")
    partie = types.Part.from_bytes(data=contenu, mime_type=mime)
    invite = _prompt(colonnes, fixes, consignes)
    derniere = None
    for modele in _CTX["modeles"]:
        for essai in range(2):
            try:
                rep = client.models.generate_content(
                    model=modele,
                    contents=[partie, invite],
                    config=types.GenerateContentConfig(
                        temperature=0, response_mime_type="application/json"),
                )
                return _parse_json(rep.text)
            except Exception as e:  # surcharge, quota, JSON illisible...
                derniere = e
                msg = str(e).lower()
                if "429" in msg or "quota" in msg or "resource_exhausted" in msg:
                    break  # inutile de réessayer ce modèle
    raise RuntimeError(f"Lecture impossible pour « {nom} » : {derniere}")


def _cle_doublon(rec, colonnes):
    nom = next((rec.get(c) for c in colonnes if _norm(c) == "nom"), "")
    pren = next((rec.get(c) for c in colonnes if _norm(c).startswith("prenom")), "")
    dn = next((rec.get(c) for c in colonnes if _role(c) == "date"), "")
    return (_norm(nom), _norm(pren), _norm(dn))


def _nettoyer(rec, colonnes, fixes):
    """Normalise un enregistrement brut de Gemini -> {'valeurs': {...}, 'doutes': {col: raison}}."""
    valeurs, doutes = {}, {}
    for c in colonnes:
        v = rec.get(c)
        v = "" if _vide(v) else str(v).strip()
        role = _role(c)
        d = None
        if role == "date":
            v, d = normaliser_date(v)
        elif role == "tel":
            v, d = normaliser_telephones(v)
        elif role == "sexe":
            v, d = normaliser_sexe(v)
        elif _norm(c) == "nom":
            v = v.upper()
        valeurs[c] = v
        if d:
            doutes[c] = d
    for c, v in fixes.items():
        if c in valeurs:
            valeurs[c] = v
    for ligne in rec.get("_doutes") or []:
        ligne = str(ligne)
        col, _, raison = ligne.partition(":")
        cible = next((c for c in colonnes if _norm(c) == _norm(col)), None)
        if cible:
            doutes[cible] = (doutes[cible] + " · " if cible in doutes else "") + raison.strip()
        else:
            doutes.setdefault(colonnes[0], ligne.strip())
    return {"valeurs": valeurs, "doutes": doutes}


def extraire_fiches(fichiers, colonnes=None, fixes=None, consignes="", existants=None, progression=None):
    """
    fichiers  : liste de (nom, bytes). Types : jpg, png, webp, heic, pdf.
    colonnes  : noms des colonnes (défaut : COLONNES_DEFAUT).
    fixes     : {colonne: valeur} appliqué à tous les élèves (ex. {"Nationalité": "Beninoise"}).
    existants : clés (nom, prénom, date) déjà présentes dans la liste -> ignorées (doublons).
    Retourne {"lignes": [{"valeurs", "doutes", "source"}], "doublons": [...], "erreurs": [...]}.
    """
    colonnes = list(colonnes or COLONNES_DEFAUT)
    fixes = {c: v for c, v in (fixes or {}).items() if c in colonnes and not _vide(v)}
    vus = set(existants or [])
    lignes, doublons, erreurs = [], [], []
    for i, (nom, contenu) in enumerate(fichiers, 1):
        ext = "." + nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
        mime = TYPES_ACCEPTES.get(ext)
        if progression:
            progression(i, len(fichiers), nom)
        if not mime:
            erreurs.append(f"{nom} : type de fichier non pris en charge")
            continue
        try:
            bruts = _lire_un_fichier(nom, contenu, mime, colonnes, fixes, consignes)
        except Exception as e:
            erreurs.append(str(e))
            continue
        if not bruts:
            erreurs.append(f"{nom} : aucune fiche trouvée")
        for brut in bruts:
            ligne = _nettoyer(brut, colonnes, fixes)
            ligne["source"] = nom
            cle = _cle_doublon(ligne["valeurs"], colonnes)
            if cle in vus and any(cle):
                doublons.append(ligne)
                continue
            vus.add(cle)
            lignes.append(ligne)
    return {"lignes": lignes, "doublons": doublons, "erreurs": erreurs, "colonnes": colonnes}


# ───────────────────────────────── Excel ─────────────────────────────────

def lire_modele(xlsx_bytes):
    """Lit un classeur existant : {feuille: {'colonnes': [...], 'cles': [(nom,prénom,date)...], 'lignes': n}}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    info = {}
    for ws in wb.worksheets:
        cols = [str(c.value).strip() for c in ws[1] if not _vide(c.value)] if ws.max_row >= 1 else []
        if not cols:
            continue
        cles, n = [], 0
        for r in range(2, ws.max_row + 1):
            rec = {cols[j]: ws.cell(r, j + 1).value for j in range(len(cols))}
            if all(_vide(v) for v in rec.values()):
                continue
            n += 1
            cles.append(_cle_doublon({k: ("" if v is None else v) for k, v in rec.items()}, cols))
        info[ws.title] = {"colonnes": cols, "cles": cles, "lignes": n}
    return info


def ecrire_xlsx(lignes, colonnes, xlsx_bytes=None, feuille=None, nom_feuille_neuve="Liste"):
    """
    Écrit les lignes {'valeurs','doutes'} :
      - à la suite de `feuille` dans le classeur fourni (mise en forme de la dernière ligne recopiée),
      - ou dans un nouveau classeur. Les cellules douteuses sont en jaune avec un commentaire.
    Retourne les octets du .xlsx.
    """
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    jaune = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    if xlsx_bytes:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb[feuille] if feuille and feuille in wb.sheetnames else wb.worksheets[0]
        entetes = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        # dernière ligne réellement remplie
        derniere = 1
        for r in range(ws.max_row, 1, -1):
            if any(not _vide(ws.cell(r, c).value) for c in range(1, len(entetes) + 1)):
                derniere = r
                break
        modele_ligne = derniere if derniere > 1 else None
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = nom_feuille_neuve
        for j, c in enumerate(colonnes, 1):
            cell = ws.cell(1, j, c)
            cell.font = Font(bold=True)
            ws.column_dimensions[get_column_letter(j)].width = max(14, len(c) + 4)
        ws.freeze_panes = "A2"
        entetes = list(colonnes)
        derniere, modele_ligne = 1, None

    index = {_norm(h): j for j, h in enumerate(entetes, 1) if h}
    r = derniere
    for ligne in lignes:
        r += 1
        for c in colonnes:
            j = index.get(_norm(c))
            if not j:
                continue
            v = ligne["valeurs"].get(c, "")
            cell = ws.cell(r, j, v if v != "" else None)
            cell.number_format = "@"  # texte : garde le 0 des numéros et les dates JJ/MM/AAAA
            if modele_ligne:
                src = ws.cell(modele_ligne, j)
                cell.font, cell.border, cell.alignment = copy(src.font), copy(src.border), copy(src.alignment)
            if c in ligne.get("doutes", {}):
                cell.fill = jaune
                cell.comment = Comment("À vérifier : " + ligne["doutes"][c], "Jarvis")
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()
