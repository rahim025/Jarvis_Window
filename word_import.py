"""
Photo (ou PDF) d'un texte écrit à la main -> texte modifiable -> document Word (.docx).

1. transcrire() : Gemini lit l'écriture et renvoie un texte avec une mise en forme très simple :
       # Titre      ## Sous-titre      - puce      1. liste numérotée
       ligne vide = nouveau paragraphe ; retour à la ligne simple = saut de ligne dans le paragraphe
   plus la liste des passages incertains.
2. L'utilisateur corrige le texte à l'écran.
3. creer_docx() : transforme ce texte en vrai fichier Word.
"""
import io
import json
import re

import fiches_import
from fiches_import import TYPES_ACCEPTES, _CTX

_PROMPT = """Tu transcris fidèlement un texte écrit à la main (photo ou PDF) en texte numérique, en français sauf si le texte est dans une autre langue.

Règles :
- Recopie EXACTEMENT ce qui est écrit, dans l'ordre de lecture. N'ajoute rien, ne résume pas, ne reformule pas.{orthographe}
- Garde la structure : un titre écrit en haut ou souligné -> ligne commençant par "# " (sous-titre : "## "). Une liste à puces ou à tirets -> lignes commençant par "- ". Une liste numérotée -> "1. ", "2. ", etc.
- Sépare les paragraphes par UNE ligne vide. Les lignes d'un même paragraphe, coupées seulement par la fin de la ligne de papier, doivent être rejointes en une seule ligne de texte. Garde un retour à la ligne simple uniquement quand l'auteur va vraiment à la ligne (adresse, vers, formule de politesse...).
- Mot ou passage illisible : écris [illisible]. Mot incertain : écris ton meilleur choix et signale-le dans "doutes".
- Mots barrés ou ratures : ignore-les.
- Si l'image contient plusieurs pages, enchaîne-les dans l'ordre.{extra}

Réponds UNIQUEMENT par un objet JSON : {{"texte": "<le texte transcrit>", "doutes": ["<passage incertain> : <raison>", ...]}}"""


def _prompt(corriger, consignes):
    ortho = ("\n- Corrige les fautes d'orthographe, d'accord et de ponctuation, sans changer le sens ni le style."
             if corriger else
             "\n- Garde l'orthographe de l'auteur, même s'il y a des fautes.")
    extra = f"\nConsigne de l'utilisateur : {consignes.strip()}" if consignes and consignes.strip() else ""
    return _PROMPT.format(orthographe=ortho, extra=extra)


def _lire_json(texte):
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", (texte or "").strip(), flags=re.S).strip()
    try:
        d = json.loads(t)
    except Exception:
        m = re.search(r"\{.*\}", t, flags=re.S)
        if not m:
            return {"texte": t, "doutes": []}  # Gemini a répondu en texte brut : on le garde tel quel
        d = json.loads(m.group(0))
    if isinstance(d, list):
        d = d[0] if d and isinstance(d[0], dict) else {"texte": "\n".join(map(str, d))}
    return {"texte": str(d.get("texte", "")).strip(), "doutes": [str(x) for x in (d.get("doutes") or [])]}


def transcrire(fichiers, corriger=False, consignes="", progression=None):
    """fichiers : liste de (nom, bytes) dans l'ordre des pages. Retourne {"texte", "doutes", "erreurs"}."""
    client, types = _CTX["client"], _CTX["types"]
    if client is None or types is None:
        raise RuntimeError("Gemini n'est pas configuré (clé GEMINI_API_KEY manquante).")
    invite = _prompt(corriger, consignes)
    textes, doutes, erreurs = [], [], []
    for i, (nom, contenu) in enumerate(fichiers, 1):
        ext = "." + nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
        mime = TYPES_ACCEPTES.get(ext)
        if progression:
            progression(i, len(fichiers), nom)
        if not mime:
            erreurs.append(f"{nom} : type de fichier non pris en charge")
            continue
        partie = types.Part.from_bytes(data=contenu, mime_type=mime)
        derniere, resultat = None, None
        for modele in _CTX["modeles"]:
            for _ in range(2):
                try:
                    rep = client.models.generate_content(
                        model=modele, contents=[partie, invite],
                        config=types.GenerateContentConfig(temperature=0, response_mime_type="application/json"))
                    resultat = _lire_json(rep.text)
                    break
                except Exception as e:
                    derniere = e
                    if any(k in str(e).lower() for k in ("429", "quota", "resource_exhausted")):
                        break
            if resultat is not None:
                break
        if resultat is None:
            erreurs.append(f"Lecture impossible pour « {nom} » : {derniere}")
            continue
        if not resultat["texte"]:
            erreurs.append(f"{nom} : aucun texte trouvé")
            continue
        textes.append(resultat["texte"])
        doutes += [(f"{nom} — " if len(fichiers) > 1 else "") + d for d in resultat["doutes"]]
    return {"texte": "\n\n".join(textes), "doutes": doutes, "erreurs": erreurs}


# ───────────────────────────────── Word ─────────────────────────────────

def creer_docx(texte, titre_fichier=""):
    """Texte (mise en forme simple, voir plus haut) -> octets d'un .docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    if titre_fichier:
        doc.core_properties.title = titre_fichier

    # blocs séparés par une ligne vide
    for bloc in re.split(r"\n\s*\n", (texte or "").replace("\r\n", "\n").strip()):
        lignes = [l.rstrip() for l in bloc.split("\n") if l.strip()]
        if not lignes:
            continue
        i = 0
        while i < len(lignes):
            l = lignes[i]
            m = re.match(r"^(#{1,3})\s+(.*)$", l)
            if m:
                doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
                i += 1
            elif re.match(r"^\s*[-•*]\s+", l):
                doc.add_paragraph(re.sub(r"^\s*[-•*]\s+", "", l), style="List Bullet")
                i += 1
            elif re.match(r"^\s*\d+[.)]\s+", l):
                doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", l), style="List Number")
                i += 1
            else:  # paragraphe : lignes consécutives = sauts de ligne dans le même paragraphe
                p = doc.add_paragraph()
                debut = True
                while i < len(lignes) and not re.match(r"^(#{1,3}\s|\s*[-•*]\s|\s*\d+[.)]\s)", lignes[i]):
                    if not debut:
                        p.add_run().add_break()
                    p.add_run(lignes[i].strip())
                    debut = False
                    i += 1
    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()
